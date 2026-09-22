"""
Generates the MVC-ARENA quantitative knowledge test (three parallel sets).

Purpose: measure whether playing the game improves understanding of what the
game actually teaches. The three aims of the game, and the item blocks that
test them, are:

  Aim 1  Separation of concerns .............. Q1-Q5
  Aim 2  Task of every component ............. Q6-Q13
  Aim 3a Special-card gameplay knowledge ..... Q14-Q15
  Aim 3b Special-card concepts and transfer .. Q16-Q20

Request/call-flow items from the earlier questionnaire are deliberately
excluded: the game does not teach the MVC request cycle.

The three sets (A, B, C) share one blueprint, so question N tests the same
category in every set. Use one set as pre-test, a different set as post-test,
and rotate the assignment across participants.

Run:    python quant_questionnaire.py   (from the supportdocs folder)
Outputs:
  MVC-ARENA_Quant_Test_Set_A.docx
  MVC-ARENA_Quant_Test_Set_B.docx
  MVC-ARENA_Quant_Test_Set_C.docx
  MVC-ARENA_Quant_Test_Researcher_Guide.docx
"""

import os
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

# ---------------------------------------------------------------------------
# Blueprint: (question number, aim, category)
# ---------------------------------------------------------------------------
AIMS = {
    1: "Separation of concerns",
    2: "Task of every component",
    3: "Special-card gameplay knowledge",
    4: "Special-card concepts and transfer",
}

BLUEPRINT = [
    (1, 1, "What MVC is (purpose of the pattern)"),
    (2, 1, "Main benefit of separating concerns"),
    (3, 1, "Change one layer without affecting the others"),
    (4, 1, "Which layer owns a given responsibility"),
    (5, 1, "Spot a misplaced responsibility"),
    (6, 2, "Task of the Model layer"),
    (7, 2, "Task of the View layer"),
    (8, 2, "Task of the Controller layer"),
    (9, 2, "Task of a specific Model component"),
    (10, 2, "Task of a specific View component"),
    (11, 2, "Task of a specific Controller component"),
    (12, 2, "Classify a component into its layer"),
    (13, 2, "Classify a component into its layer"),
    (14, 3, "Game mechanic: Inheritance card"),
    (15, 3, "Game mechanic: Polymorphism card"),
    (16, 4, "Concept: inheritance"),
    (17, 4, "Concept: interface"),
    (18, 4, "Apply a defensive component to a security scenario"),
    (19, 4, "Apply logging to defect diagnosis"),
    (20, 4, "Apply version control and backups to recovery"),
]

# ---------------------------------------------------------------------------
# Question sets.  Each item: (stem, [options a-d], index_of_correct_option)
# Answer keys are derived from the index so they cannot drift out of sync.
# Component-to-layer and attack-to-defense answers follow the game's own
# classification (src/classes/card/cardDescriptions.js, cardData.js).
# ---------------------------------------------------------------------------
SETS = {
    "SET A": [
        # --- Aim 1: separation of concerns ---
        ("What is the main purpose of the MVC pattern?",
         ["to make an application run faster",
          "to organize an application into Model, View, and Controller, each with a separate responsibility",
          "to encrypt user data",
          "to replace the database"], 1),
        ("The main benefit of separating concerns in MVC is that:",
         ["each part can be changed without breaking the others",
          "the program needs fewer files",
          "all bugs are removed automatically",
          "the application never needs testing"], 0),
        ("You want to redesign the user interface without touching any data or business rules. "
         "In MVC you mainly change the:",
         ["Model", "Controller", "View", "Database"], 2),
        ("Which layer should own the rule \"an order total equals the sum of its line items\"?",
         ["View", "Model", "Controller", "Routing"], 1),
        ("A developer writes database queries directly inside the screen-drawing code. "
         "This breaks separation of concerns because:",
         ["the View is now doing the Model's job",
          "the Model is now choosing the screen layout",
          "the Controller is now presenting the interface",
          "the database is now routing user requests"], 0),

        # --- Aim 2: task of every component ---
        ("The Model layer is mainly responsible for:",
         ["drawing buttons and pages",
          "the application's data and business logic",
          "mapping URLs to code",
          "styling the interface"], 1),
        ("The View layer is mainly responsible for:",
         ["storing data",
          "checking permissions",
          "presenting information to the user",
          "writing to the database"], 2),
        ("The Controller layer is mainly responsible for:",
         ["receiving requests and coordinating the Model and View",
          "saving files to disk",
          "displaying the interface",
          "encrypting passwords"], 0),
        ("What does an ORM (Object-Relational Mapping) component do?",
         ["displays database tables on screen",
          "maps database tables to code objects and builds safe, parameterized queries",
          "limits how many requests a user can send",
          "verifies who a user is"], 1),
        ("What does an Output Validation component do?",
         ["escapes and encodes data right before it is displayed so embedded scripts cannot run",
          "stores frequently used data in fast memory",
          "routes requests to the right handler",
          "tracks code history"], 0),
        ("What does a Routing component do?",
         ["stores passwords securely",
          "draws the mobile interface",
          "maps incoming URLs and requests to the code that should handle them",
          "saves files to the cloud"], 2),
        ("In MVC-ARENA, a Caching component belongs to which lane?",
         ["View", "Controller", "Model", "None of these"], 2),
        ("In MVC-ARENA, an Authentication component belongs to which lane?",
         ["Model", "View", "Controller", "Database"], 2),

        # --- Aim 3a: game mechanics ---
        ("In MVC-ARENA, what happens when an Inheritance card is played?",
         ["It starts a new stack in the selected MVC lane",
          "It doubles the value of an existing stack",
          "It changes into any component in that lane",
          "It removes a destructive-event card"], 0),
        ("In MVC-ARENA, what happens when a Polymorphism card is played in an MVC lane?",
         ["It starts an empty stack in another lane",
          "It becomes a selected component from that lane",
          "It doubles every stack in the selected lane",
          "It blocks the next attack against that lane"], 1),

        # --- Aim 3b: concepts and transfer ---
        ("A specialized class reuses common behavior from a more general parent class. "
         "Which concept does this demonstrate?",
         ["Interface", "Logging", "Inheritance", "Rate limiting"], 2),
        ("Several storage classes can be exchanged because each promises the same set of operations. "
         "What enables this?",
         ["A shared interface", "A database backup", "An error log", "A request limit"], 0),
        ("An application builds database commands from untrusted text. Which practice most directly "
         "reduces the risk of malicious query injection?",
         ["Escaping screen output", "Limiting request frequency",
          "Using parameterized queries through an ORM", "Recording errors in a log"], 2),
        ("A defect occurs only occasionally in production. What information would best help a "
         "developer reconstruct what happened?",
         ["A history of recorded events and errors", "A list of screen colours",
          "A copy of the navigation menu", "A count of database tables"], 0),
        ("A faulty release corrupts important project files. Which preparation best supports recovery?",
         ["Keeping version history and tested backups", "Adding another screen layout",
          "Increasing the request limit", "Renaming the controller classes"], 0),
    ],

    "SET B": [
        # --- Aim 1 ---
        ("MVC is best described as:",
         ["a programming language",
          "a type of database",
          "a design pattern that splits an application into three parts with distinct roles",
          "a security certificate"], 2),
        ("Because concerns are separated in MVC, a team can:",
         ["skip writing tests",
          "work on the interface and the data logic independently",
          "keep the whole app in one file",
          "avoid all security threats"], 1),
        ("You must switch from a web page to a mobile app but keep all the same data and rules. "
         "In MVC you replace the:",
         ["View", "Model", "Controller", "Database"], 0),
        ("Deciding which code should handle the URL \"/login\" belongs to the:",
         ["Model", "View", "Controller", "Cache"], 2),
        ("A developer writes the \"calculate order total\" logic inside the request-handling code. "
         "Which layer should own that logic instead?",
         ["View", "Model", "Controller", "Database"], 1),

        # --- Aim 2 ---
        ("Business rules and stored application data live mainly in the:",
         ["Controller", "View", "Middleware", "Model"], 3),
        ("The part of the application the user sees and interacts with is the:",
         ["Model", "View", "Controller", "ORM"], 1),
        ("Handling an incoming request and deciding what to do with it is the job of the:",
         ["Controller", "Model", "View", "Cache"], 0),
        ("What does a Secrets Manager component do?",
         ["draws the login page",
          "securely stores and controls access to passwords and API keys",
          "routes requests",
          "limits traffic"], 1),
        ("What does a Web View component do?",
         ["stores data in the database",
          "checks user permissions",
          "provides the browser-based user interface",
          "maps tables to objects"], 2),
        ("What does a Rate Limiting component do?",
         ["caps how many requests a client can make in a time window",
          "encrypts the database",
          "displays error pages",
          "tracks code history"], 0),
        ("In MVC-ARENA, a Mobile View component belongs to which lane?",
         ["Model", "Controller", "View", "Storage"], 2),
        ("In MVC-ARENA, an ORM component belongs to which lane?",
         ["View", "Model", "Controller", "None of these"], 1),

        # --- Aim 3a: game mechanics ---
        ("In MVC-ARENA, where can an Inheritance card be placed?",
         ["Only on an opponent's completed stack",
          "In an MVC lane to begin a new stack",
          "Only on top of a defensive card",
          "In the redraw pile to gain another turn"], 1),
        ("In MVC-ARENA, after placing Polymorphism in the View lane, what does the player do?",
         ["Selects which View component the card becomes",
          "Moves the card into the Model lane",
          "Doubles every View stack automatically",
          "Discards one of the opponent's cards"], 0),

        # --- Aim 3b: concepts and transfer ---
        ("A new report class extends a general document class and keeps its shared behaviour. "
         "Which concept is being used?",
         ["Polymorphism", "Inheritance", "Logging", "Routing"], 1),
        ("Payment services from different providers can replace one another because they expose "
         "the same required operations. What defines those operations?",
         ["A cache", "A backup", "An interface", "A log"], 2),
        ("A public service becomes unavailable because one client sends excessive traffic. "
         "Which practice most directly limits this behaviour?",
         ["Encoding displayed output", "Applying rate limits to requests",
          "Recording application errors", "Mapping tables to objects"], 1),
        ("A program fails after an unusual sequence of actions. Which evidence would be most useful "
         "for tracing the cause?",
         ["Saved event and error logs", "The interface font settings",
          "The number of view classes", "The names of menu buttons"], 0),
        ("An update introduces serious faults and must be undone. Which preparation is most useful?",
         ["A version-control history with a known working revision", "A larger cache",
          "A different web-page theme", "A higher traffic allowance"], 0),
    ],

    "SET C": [
        # --- Aim 1 ---
        ("Which statement about MVC is correct?",
         ["it stores everything in one large file",
          "it divides an application into Model, View, and Controller, each with its own responsibility",
          "it is only used for games",
          "it removes the need for a user interface"], 1),
        ("\"Separation of concerns\" means:",
         ["every part of the program does everything",
          "each part has one clear responsibility and does not do the others' jobs",
          "the application is split across many servers",
          "concerns are hidden from the user"], 1),
        ("The database engine changes from MySQL to PostgreSQL. With good MVC separation, "
         "which layer is mostly affected?",
         ["View", "Controller", "Model", "all three equally"], 2),
        ("Showing an error message nicely on the screen is the responsibility of the:",
         ["Model", "View", "Controller", "Database"], 1),
        ("A View component directly changes a user's account balance in the database. What is wrong?",
         ["the View is doing the Model's data-logic job",
          "the Model is doing the View's presentation job",
          "the Controller is doing the View's presentation job",
          "the database is doing the Controller's routing job"], 0),

        # --- Aim 2 ---
        ("Where would you put logic that validates and saves a user record to storage?",
         ["View", "Controller", "Model", "Router"], 2),
        ("A command-line interface that shows output to the user is part of the:",
         ["View", "Model", "Controller", "Cache"], 0),
        ("Middleware that checks every request before it reaches the main logic belongs to the:",
         ["Model", "View", "Database", "Controller"], 3),
        ("What does a Caching component do?",
         ["escapes output for display",
          "stores frequently used data in fast memory to speed up responses",
          "verifies who a user is",
          "maps URLs to handlers"], 1),
        ("What does a CLI View component do?",
         ["stores files on disk",
          "lets the user interact with the app through text commands",
          "limits requests",
          "tracks code history"], 1),
        ("What does a Middleware component do?",
         ["runs between a request and its handler for tasks like auth checks or parsing",
          "draws the mobile interface",
          "stores passwords",
          "saves files to the cloud"], 0),
        ("In MVC-ARENA, an Output Validation component belongs to which lane?",
         ["Model", "Controller", "View", "Storage"], 2),
        ("In MVC-ARENA, an Authorization component belongs to which lane?",
         ["View", "Model", "Database", "Controller"], 3),

        # --- Aim 3a: game mechanics ---
        ("In MVC-ARENA, which card is used to begin a new stack in a Model, View, or Controller lane?",
         ["Logger", "Polymorphism", "Inheritance", "Error Handling"], 2),
        ("In MVC-ARENA, how is the component represented by a Polymorphism card determined?",
         ["It always becomes a Database component",
          "It copies the opponent's last component",
          "It is selected from the lane where the card was placed",
          "It is chosen randomly from all three lanes"], 2),

        # --- Aim 3b: concepts and transfer ---
        ("A mobile screen class extends a general screen class instead of repeating its common code. "
         "Which concept does this demonstrate?",
         ["Logging", "Inheritance", "Rate limiting", "Output validation"], 1),
        ("Different notification services can be swapped without changing their caller because all "
         "provide the same required methods. What specifies those methods?",
         ["An interface", "A log file", "A cache entry", "A backup"], 0),
        ("Untrusted text is about to be displayed on a web page. Which practice most directly "
         "prevents embedded scripts from running?",
         ["Increasing the cache size", "Encoding or validating the output",
          "Saving a version-control revision", "Recording the request in a log"], 1),
        ("Users report an error that developers cannot reproduce. What should developers inspect first?",
         ["Recorded application events and errors", "The colours used by the interface",
          "The order of navigation links", "The number of model classes"], 0),
        ("A developer accidentally deletes working code. Which preparation most directly enables restoration?",
         ["A higher request limit", "A different command-line view",
          "A version-control history and backup", "An additional routing rule"], 2),
    ],
}

LETTERS = "abcd"

# Use exactly five correct answers in each position on every form. Questions
# are authored for content first, then their options are deterministically
# repositioned here. Different forms use different patterns to avoid a shared
# answer sequence.
TARGET_KEYS = {
    "SET A": [0, 1, 2, 3, 1, 2, 3, 0, 2, 3, 0, 1, 3, 0, 1, 2, 3, 0, 1, 2],
    "SET B": [1, 2, 3, 0, 2, 3, 0, 1, 3, 0, 1, 2, 0, 1, 2, 3, 0, 1, 2, 3],
    "SET C": [2, 3, 0, 1, 3, 0, 1, 2, 0, 1, 2, 3, 1, 2, 3, 0, 1, 2, 3, 0],
}


def position_correct_answer(item, target):
    """Move the correct option to target while preserving distractor order."""
    stem, options, correct = item
    correct_option = options[correct]
    distractors = [option for index, option in enumerate(options) if index != correct]
    reordered = distractors[:]
    reordered.insert(target, correct_option)
    return stem, reordered, target


SETS = {
    set_name: [
        position_correct_answer(item, target)
        for item, target in zip(items, TARGET_KEYS[set_name])
    ]
    for set_name, items in SETS.items()
}

# ---------------------------------------------------------------------------
# Sanity checks: every set must match the blueprint and have a valid key.
# ---------------------------------------------------------------------------
for set_name, items in SETS.items():
    assert len(items) == len(BLUEPRINT), f"{set_name} has {len(items)} items, expected {len(BLUEPRINT)}"
    for n, (stem, options, correct) in enumerate(items, start=1):
        assert len(options) == 4, f"{set_name} Q{n} must have 4 options"
        assert 0 <= correct < 4, f"{set_name} Q{n} has invalid correct index {correct}"
    counts = [sum(item[2] == index for item in items) for index in range(4)]
    assert counts == [5, 5, 5, 5], f"{set_name} answer positions are unbalanced: {counts}"

# ---------------------------------------------------------------------------
# Document helpers
# ---------------------------------------------------------------------------
doc = Document()
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)


def heading(text, level=1):
    return doc.add_heading(text, level=level)


def para(text="", bold=False, italic=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    return p


def bullet(text):
    doc.add_paragraph(text, style="List Bullet")


def mcq(num, stem, options):
    p = doc.add_paragraph()
    p.add_run(f"{num}. ").bold = True
    p.add_run(stem)
    for i, opt in enumerate(options):
        op = doc.add_paragraph()
        op.paragraph_format.left_indent = Pt(24)
        op.add_run(f"{LETTERS[i]}) ").bold = True
        op.add_run(opt)


def table(headers, rows, style="Light Grid Accent 1"):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.style = style
    for i, h in enumerate(headers):
        tbl.rows[0].cells[i].paragraphs[0].add_run(h).bold = True
    for row in rows:
        cells = tbl.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
    return tbl


# ---------------------------------------------------------------------------
# Title + administration
# ---------------------------------------------------------------------------
title = doc.add_heading("MVC-ARENA Quantitative Knowledge Test", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

para(
    "This test measures understanding of the three things the game is designed to "
    "teach: separation of concerns, the task of every MVC component, and the use "
    "and activity of the special cards. It comes in three parallel sets (A, B, C) "
    "that share one blueprint, so question N tests the same category in every set.",
    italic=True,
)

heading("How to administer", level=2)
bullet("Give one set before playing (pre-test) and a different set after (post-test). "
       "Optionally give the third set later as a retention test.")
bullet("For pre/post testing, assign participants as evenly as possible across the six "
       "ordered pairs: A-B, B-A, A-C, C-A, B-C, and C-B.")
bullet("Allow about 15 minutes per set. No notes or access to the game during the test.")

heading("Scoring", level=2)
bullet("1 point per correct answer. Total score: 0-20.")
bullet("Score a blank or an item with multiple marked answers as incorrect.")
bullet("Do not interpret raw post-test minus pre-test scores without accounting for "
       "which form was administered. Include test form and pre-test score in the analysis.")
bullet("Report four sub-scores. Treat gameplay knowledge as a manipulation check, "
       "not by itself as evidence of increased conceptual understanding:")
table(
    ["Aim", "Questions", "Max"],
    [
        ("1. Separation of concerns", "Q1 - Q5", 5),
        ("2. Task of every component", "Q6 - Q13", 8),
        ("3a. Special-card gameplay knowledge", "Q14 - Q15", 2),
        ("3b. Special-card concepts and transfer", "Q16 - Q20", 5),
    ],
)
para()
para(
    "Request / call-flow items are intentionally excluded: the game does not teach "
    "the MVC request cycle, so testing it would not reflect the game's learning aims.",
    italic=True,
)

# ---------------------------------------------------------------------------
# Blueprint table
# ---------------------------------------------------------------------------
heading("Blueprint (same for Sets A, B, C)", level=2)
table(
    ["Q", "Aim", "Category"],
    [(q, AIMS[aim], cat) for q, aim, cat in BLUEPRINT],
)

# ---------------------------------------------------------------------------
# The three sets
# ---------------------------------------------------------------------------
for set_name, items in SETS.items():
    doc.add_page_break()
    heading(set_name, level=1)
    para("Participant ID: ______________    Date: ____________    "
         "(circle one: pre-test / post-test / retention)", italic=True)
    para("Choose the single best answer for each question.", italic=True)

    current_aim = None
    for n, (stem, options, _correct) in enumerate(items, start=1):
        aim = BLUEPRINT[n - 1][1]
        if aim != current_aim:
            current_aim = aim
            heading(f"Part {aim} - {AIMS[aim]}", level=3)
        mcq(n, stem, options)

# ---------------------------------------------------------------------------
# Answer keys (derived from the correct index)
# ---------------------------------------------------------------------------
doc.add_page_break()
heading("Answer Keys", level=1)
for set_name, items in SETS.items():
    key = ", ".join(f"{n}-{LETTERS[correct]}" for n, (_s, _o, correct) in enumerate(items, start=1))
    p = doc.add_paragraph()
    p.add_run(f"{set_name}: ").bold = True
    p.add_run(key)

para()
heading("Per-aim key", level=2)
rows = []
for set_name, items in SETS.items():
    for aim_id in AIMS:
        qs = [n for n, _aim, _cat in ((q, a, c) for q, a, c in BLUEPRINT) if _aim == aim_id]
        key = ", ".join(f"{n}-{LETTERS[items[n - 1][2]]}" for n in qs)
        rows.append((set_name, AIMS[aim_id], key))
table(["Set", "Aim", "Key"], rows)

para()
note = doc.add_paragraph()
note.add_run("Note on classification: ").bold = True
note.add_run(
    "Component-to-layer and attack-to-defense answers follow the game's own rules "
    "(e.g. Authentication, Authorization, Rate Limiting, Routing and Middleware are "
    "Controller; Output Validation is View; Caching, ORM, Data Validation and Secrets "
    "Manager are Model). This is correct for measuring what the game teaches. Note in "
    "your methodology that some of these can sit in other layers in different "
    "real-world designs."
)

output_dir = os.path.dirname(os.path.abspath(__file__))
guide_out = os.path.join(output_dir, "MVC-ARENA_Quant_Test_Researcher_Guide.docx")
doc.save(guide_out)
print("wrote", guide_out)

# Produce one clean participant file per form. These files intentionally omit
# the blueprint, aim labels, scoring guidance, and answer keys.
for set_name, items in SETS.items():
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    title = doc.add_heading(f"MVC-ARENA Knowledge Test - {set_name}", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para("Participant ID: ______________    Date: ____________    "
         "(circle one: pre-test / post-test / retention)", italic=True)
    para("Choose the single best answer for each question. Mark only one answer per question.",
         italic=True)

    for n, (stem, options, _correct) in enumerate(items, start=1):
        mcq(n, stem, options)

    form_letter = set_name[-1]
    participant_out = os.path.join(
        output_dir, f"MVC-ARENA_Quant_Test_Set_{form_letter}.docx"
    )
    doc.save(participant_out)
    print("wrote", participant_out)
